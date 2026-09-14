# -*- coding: utf-8 -*-
"""生成参数实验用的测试文档：格式（PDF/DOCX/MD/TXT）、长度、语言各异。
每份文档埋入唯一数值型「锚点事实」，答案可程序化验证；
doc5/doc6 是同构干扰对，专门测检索区分度。"""
import os
import random

import fitz
import docx

OUT = os.path.join(os.path.dirname(__file__), "docs")
os.makedirs(OUT, exist_ok=True)
random.seed(20260829)

# ---------- doc1: 中文长文档，PDF 多页 ----------
QD_SECTIONS = [
    ("一、技术背景与路线", [
        "量子点显示（QLED）利用胶体量子点的窄谱发射特性，是下一代显示技术的核心路线之一。"
        "与传统 OLED 相比，量子点在色纯度、寿命与亮度维持方面具有理论优势。",
        "本报告汇总实验室 2024 至 2026 年度三个研究阶段的成果，覆盖材料合成、器件架构与量产工艺三个方向。",
        "红光体系以 CdSe-9 为核心材料，其光谱半峰宽收窄至 18 纳米，达到行业领先的色域覆盖水平。",
        "绿光体系沿用 InP 基方案，外量子效率稳定在 21.2%，主要瓶颈仍是器件稳定性。",
    ]),
    ("二、蓝光器件的关键突破", [
        "蓝光是量子点显示最难攻克的颜色窗口。研究团队通过配体交换工艺与核壳结构优化，"
        "开发出代号为 ZnSe-77 的新型蓝光材料。",
        "在 1000 尼特初始亮度条件下，ZnSe-77 器件的外量子效率达到 23.7%，较上一代提升 3.5 个百分点。",
        "寿命方面，器件的 T95 寿命达到 41000 小时，满足电视整机五年的亮度维持要求。",
        "蓝光器件的光谱半峰宽为 21 纳米，色坐标偏离 Rec.2020 标准蓝点小于 0.004。",
    ]),
    ("三、材料合成与工艺", [
        "ZnSe-77 的合成采用高温热注入法，核心成核温度控制在 285 摄氏度，生长阶段梯度降温至 240 摄氏度。",
        "配体采用改进的双齿羧酸体系，纳米晶尺寸分布的标准差降至 4.2%。",
        "空穴传输层的迁移率提升至 12.4 平方厘米每伏秒，是器件效率提升的关键因素之一。",
        "所有合成实验在氮气手套箱内完成，水氧含量均低于 0.1 ppm。",
    ]),
    ("四、量产工艺与良率", [
        "2026 年第二季度，中试线完成 1500 片 6 代基板的喷墨打印验证，量产良率达到 91.3%。",
        "墨水的黏度控制在 10.5 厘泊，基板预热温度 55 摄氏度，有效抑制了咖啡环效应。",
        "量产成本测算显示，65 英寸 4K 面板的量子点材料成本较 OLED 方案低 17.8%。",
        "下一阶段的重点是 8.5 代线的转移，预计 2027 年完成设备验证。",
    ]),
    ("五、结论与展望", [
        "量子点显示在效率、寿命与成本三个维度均已具备替代 OLED 的条件。",
        "红光 CdSe-9、蓝光 ZnSe-77 与绿光 InP 体系共同构成了全彩化方案的材料基础。",
        "剩余的挑战集中在喷墨打印的均匀性控制与含镉材料的法规合规两个问题上。",
    ]),
]

# ---------- doc2: 英文长文档，TXT（压测 600 字符对英文的切片粒度） ----------
MOE_PARAS = [
    ("Introduction", [
        "Mixture-of-Experts (MoE) architectures have emerged as the dominant approach for scaling language models beyond the practical limits of dense transformers. By conditionally activating a small subset of parameters for each token, MoE models decouple model capacity from computational cost.",
        "This survey systematically reviews 48 MoE papers published between 2023 and 2026, covering routing algorithms, expert specialization analysis, training stability techniques, and deployment strategies. We propose the hypothetical MoE-X architecture family as a running example to quantify the claims discussed across the literature.",
    ]),
    ("Routing Algorithms", [
        "The MoE-X router selects the top-2 experts out of 64 available candidates for every token, a configuration that the ablation studies identify as the best trade-off between specialization and load balancing. Auxiliary load-balancing loss is weighted at 0.01.",
        "Compared with top-1 routing, the top-2 configuration improves downstream accuracy by 2.4 points on average across nine benchmarks, at the cost of a 22% increase in inference FLOPs. Router z-loss stabilization proves critical beyond the 32-expert scale.",
        "Expert capacity factors between 1.1 and 1.25 eliminate token dropping entirely on our workloads. Dropless variants based on token permutation trade a small amount of throughput for simplified deployment semantics.",
    ]),
    ("Scaling Behavior", [
        "The flagship MoE-X model is pretrained on 1.8 trillion tokens of filtered web text, code, and multilingual corpora. Total activated parameters per token are 13 billion out of 220 billion total parameters.",
        "On the MMLU benchmark suite, MoE-X closes the remaining 3.1% accuracy gap to a dense reference model with four times the activated parameter count, confirming the compute-efficiency thesis of sparse scaling.",
        "Scaling law fitting across seven model sizes yields an exponent of 0.47 for loss versus activated parameters, notably steeper than the 0.34 observed for dense models under identical data budgets.",
    ]),
    ("Serving and Efficiency", [
        "In production serving, expert parallelism across 16 nodes reduces energy consumption per query by 41% relative to the dense baseline at matched quality. The energy savings come primarily from eliminating activation recomputation.",
        "Tail latency remains the weakest point of MoE serving. The measured p99 end-to-end latency is 87 milliseconds under a 40-token generation load, driven by all-to-all communication overhead during expert dispatch.",
        "Batching strategies interact strongly with routing: dynamic batch splitting at expert boundaries recovers 31% of the throughput lost to load imbalance at high concurrency.",
    ]),
    ("Discussion and Open Problems", [
        "Expert specialization remains partially entangled: probing shows that individual experts encode overlapping linguistic features rather than clean domain boundaries. Interpretability of routing decisions is an open research area.",
        "Training instability at scale manifests as router collapse in the first 2% of training steps. The survey collects five mitigation techniques and compares their reproducibility across three independent replications.",
        "We conclude that sparse scaling is the most credible path toward the next quality tier, and identify routing robustness under distribution shift as the highest-value open problem for the community.",
    ]),
]

def english_long_text():
    parts = []
    filler = (
        "Additional analysis in the appendix replicates the headline numbers under three random seeds and "
        "reports the full deviation tables. The experimental protocol follows the standard evaluation harness "
        "released with the public benchmark, and all hyperparameters are listed in the configuration index. "
        "We additionally release the routing traces so that follow-up work can reproduce the load statistics. "
    )
    for title, paras in MOE_PARAS:
        parts.append(f"{title.upper()}\n")
        for p in paras:
            parts.append(p + " " + filler * 2 + "\n")
    return "\n".join(parts)

# ---------- doc3: 中文中等长度，DOCX，数字密集 ----------
PEROVSKITE = [
    ("实验目的", [
        "本批次实验系统考察有机无机混合钙钛矿薄膜的结晶工艺窗口，为后续模组放大提供工艺参数依据。",
        "实验围绕前驱体配比、退火曲线与反溶剂滴加时机三个变量展开，共设计 24 组对照。",
    ]),
    ("前驱体配制", [
        "标准配方中 PbI2 与 FAI 的前驱体摩尔比固定为 1.32 : 1，溶剂为 DMF 与 DMSO 的四比一混合体系。",
        "前驱体溶液在 60 摄氏度下搅拌 2 小时后经 0.22 微米 PTFE 滤头过滤，静置脱泡 30 分钟。",
        "添加 0.35 摩尔百分比的 MACl 添加剂以改善薄膜的横向结晶取向。",
    ]),
    ("薄膜制备与退火", [
        "旋涂转速为 4000 转每分钟，加速时间 3 秒，旋涂进行到第 18 秒时滴加氯苯反溶剂。",
        "薄膜在 105 摄氏度下退火 23 分钟，升温速率统一为每秒 5 摄氏度。",
        "退火气氛对结晶质量影响显著：氮气氛围样品的晶粒尺寸比干燥空气氛围大 22%。",
    ]),
    ("器件性能", [
        "冠军器件的光电转换效率达到 25.1%，开路电压 1.18 伏，短路电流密度 26.2 毫安每平方厘米。",
        "批次 B-17 的深能级缺陷密度经热导纳谱测得为 3.2 乘以 10 的 15 次方每立方厘米，是所有批次中最低的。",
        "未封装器件在相对湿度 30% 环境下放置 500 小时后仍保持初始效率的 87%。",
    ]),
    ("问题与改进", [
        "本批次出现两片基板的薄膜针孔缺陷，追溯记录显示与旋涂湿度尖峰（当日 15 时湿度 62%）相关。",
        "下一批次的改进措施包括加装除湿联锁与将反溶剂滴加时机从 18 秒提前至 15 秒。",
    ]),
]

# ---------- doc4: 中文短文档，Markdown ----------
IR_LECTURE = """# 信息检索概论 · 第七讲讲义

## 倒排索引

倒排索引是搜索引擎的核心数据结构：对词表中每个词项维护一个出现该词的文档编号列表。
构建流程依次为：分词、归一化、词项排序、合并Posting。

## 布尔检索与向量空间模型

布尔检索支持 AND / OR / NOT 的精确集合运算，但无法表达相关性强弱。
向量空间模型把文档与查询都映射为高维向量，用余弦相似度排序。

## 课程考核

本课程期末闭卷考试占综合成绩的 40%，平时作业共布置 3 次，每次占 15%，课堂出勤占 15%。
期末考试的题型为：概念辨析 4 题、计算 2 题、开放论述 1 题。
"""

# ---------- doc5 / doc6: 同构干扰对（相似句式，不同数值） ----------
def team_report(team: str, score: str, pass_rate: str, case: str, date: str) -> str:
    return f"""# {team} 自动评测报告

## 评测设置

本报告对 {team} 组研发的对话系统进行标准化评测。评测数据集为 GSM8K 数学应用题集与
内部构建的 {case} 案例集，评测日期为 {date}。

## 主要结果

在 GSM8K 测试集上，{team} 组的系统取得 {score} 分的准确率，解题过程的步骤通过率为 {pass_rate}。
与上一季度相比，主要提升来自错误反馈回路的引入。

在 {case} 案例集上，系统的一轮解决率为 71.2%，多轮收敛后提升至 86.5%。

## 失败分析

失败案例集中在多步算术与单位换算两类，合计占全部失败样例的 63.7%。
"""

TEAMA = team_report("A", "88.4", "94.2%", "医疗问诊", "2026-06-30")
TEAMB = team_report("B", "84.9", "91.8%", "法律咨询", "2026-06-30")

# ---------- 渲染 ----------
def make_pdf(path, title, sections):
    """内置 CID 字体（china-ss 宋体）+ 手动换行，产出带文本层的小体积 PDF。"""
    doc = fitz.open()

    def wrap(text, width=40):
        lines = []
        while text:
            lines.append(text[:width])
            text = text[width:]
        return lines

    for sec_title, paras in sections:
        page = doc.new_page()
        y = 90
        page.insert_text((72, 60), title, fontsize=18, fontname="china-ss")
        for p in paras:
            if y > 740:
                page = doc.new_page()
                y = 90
            page.insert_text((72, y), sec_title, fontsize=14, fontname="china-ss")
            y += 26
            for line in wrap(p):
                if y > 780:
                    page = doc.new_page()
                    y = 90
                page.insert_text((72, y), line, fontsize=11, fontname="china-ss")
                y += 18
            y += 10
    doc.save(path, garbage=4, deflate=True)
    doc.close()

def make_docx(path, sections):
    d = docx.Document()
    for sec_title, paras in sections:
        d.add_heading(sec_title, level=2)
        for p in paras:
            d.add_paragraph(p)
    d.save(path)

make_pdf(os.path.join(OUT, "量子点显示技术报告.pdf"), "量子点显示技术年度报告", QD_SECTIONS)
make_docx(os.path.join(OUT, "钙钛矿薄膜实验记录.docx"), PEROVSKITE)

with open(os.path.join(OUT, "MoE_Survey_2026.txt"), "w", encoding="utf-8") as f:
    f.write(english_long_text())

with open(os.path.join(OUT, "信息检索课程讲义.md"), "w", encoding="utf-8") as f:
    f.write(IR_LECTURE)

with open(os.path.join(OUT, "TeamA_评测报告.md"), "w", encoding="utf-8") as f:
    f.write(TEAMA)

with open(os.path.join(OUT, "TeamB_评测报告.md"), "w", encoding="utf-8") as f:
    f.write(TEAMB)

for fn in sorted(os.listdir(OUT)):
    p = os.path.join(OUT, fn)
    print(f"{fn:40s} {os.path.getsize(p):>8} bytes")
print("done ->", OUT)
